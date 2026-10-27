from datetime import datetime
from io import BytesIO
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

from .models import ScanResult


def results_frame(results: list[ScanResult]) -> pd.DataFrame:
    columns = [
        "Ürün",
        "Kondisyon",
        "Miniöde Fiyatı",
        "Piyasa Referansı",
        "Referans Kaynak",
        "Eşleşen Ürün",
        "Akakçe Fiyatı",
        "Cimri Fiyatı",
        "Akakçe Durumu",
        "Cimri Durumu",
        "Fark TL",
        "Fark %",
        "Durum",
        "Öneri",
        "SKU",
        "Barkod",
        "Ürün URL",
        "Piyasa URL",
        "Akakçe Arama",
        "Cimri Arama",
        "Arama Sorgusu",
    ]
    return pd.DataFrame([result.as_row() for result in results], columns=columns)


def create_excel(results: list[ScanResult]) -> bytes:
    frame = results_frame(results)
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        frame.to_excel(writer, index=False, sheet_name="Fiyat Analizi")
        worksheet = writer.sheets["Fiyat Analizi"]
        header_fill = PatternFill("solid", fgColor="1E3A8A")
        header_font = Font(color="FFFFFF", bold=True)
        status_column = frame.columns.get_loc("Durum") + 1

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        colors = {
            "Kritik": "FECACA",
            "Yüksek": "FED7AA",
            "Rekabetçi": "DCFCE7",
            "Dengeli": "DCFCE7",
            "Bulunamadı": "E2E8F0",
        }
        for row_index in range(2, len(frame) + 2):
            status_cell = worksheet.cell(row=row_index, column=status_column)
            for keyword, color in colors.items():
                if keyword in str(status_cell.value or ""):
                    status_cell.fill = PatternFill("solid", fgColor=color)
                    break

        for column_cells in worksheet.columns:
            letter = column_cells[0].column_letter
            max_length = max(len(str(cell.value or "")) for cell in column_cells)
            worksheet.column_dimensions[letter].width = min(max(max_length + 2, 12), 48)
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
    return buffer.getvalue()


def save_excel(report_dir: Path, content: bytes, prefix: str = "fiyat_raporu") -> Path:
    report_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = report_dir / f"{prefix}_{timestamp}.xlsx"
    path.write_bytes(content)
    return path
